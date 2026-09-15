"""DOCX Document parser for WebAssetify.

Extracts visible text, table contents, and resolves internal XML hyperlink relationships
(w:hyperlink) using python-docx.
"""

from __future__ import annotations

import io
import logging
from pathlib import Path
from typing import BinaryIO, Union

import docx
from docx.opc.constants import RELATIONSHIP_TYPE

from parsers.universal import clean_url, extract_urls

logger = logging.getLogger(__name__)

HYPERLINK_RELTYPE = RELATIONSHIP_TYPE.HYPERLINK


def extract_urls_from_docx(source: Union[str, Path, bytes, BinaryIO]) -> list[str]:
    """Extract URLs from a DOCX document (text, tables, and internal XML hyperlinks).

    Args:
        source: File path, Path object, raw bytes, or file-like binary stream.

    Returns:
        Deduplicated list of extracted URLs.
    """
    stream: BinaryIO
    if isinstance(source, (str, Path)):
        file_path = Path(source)
        if not file_path.exists():
            logger.error("DOCX file does not exist: %s", source)
            return []
        stream = io.BytesIO(file_path.read_bytes())
    elif isinstance(source, bytes):
        stream = io.BytesIO(source)
    else:
        stream = source

    urls: list[str] = []

    try:
        doc = docx.Document(stream)

        # 1. Extract URLs from document relationships (w:hyperlink targets)
        try:
            for rel in doc.part.rels.values():
                if rel.reltype == HYPERLINK_RELTYPE:
                    target = rel.target_ref
                    if target and (target.startswith("http://") or target.startswith("https://")):
                        cleaned = clean_url(target)
                        if cleaned:
                            urls.append(cleaned)
        except Exception as e:
            logger.warning("Failed to extract relationship hyperlinks: %s", e)

        # 2. Extract visible text from paragraphs
        full_text_parts: list[str] = []
        for paragraph in doc.paragraphs:
            if paragraph.text:
                full_text_parts.append(paragraph.text)

        # 3. Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        full_text_parts.append(cell.text)

        # Run universal extractor on all visible text
        if full_text_parts:
            combined_text = "\n".join(full_text_parts)
            urls.extend(extract_urls(combined_text))

    except Exception as e:
        logger.error("Error reading DOCX document: %s", e)

    # Deduplicate while preserving order
    seen: set[str] = set()
    deduped: list[str] = []
    for u in urls:
        if u not in seen:
            seen.add(u)
            deduped.append(u)

    return deduped
